import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_contents_and_details(filepath):
    print("=" * 70)
    print(f"ANALYSIS OF CONTENT & COMPLIANCE: {filepath}")
    print("=" * 70)
    doc = docx.Document(filepath)

    # 1. Title page check
    print("--- 1. TITLE PAGE ---")
    for i in range(min(15, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        txt = p.text.strip()
        if txt:
            runs_info = [(r.text[:30], r.font.size.pt if r.font.size else "default", r.bold) for r in p.runs]
            print(f"  P#{i}: '{txt[:70]}' | Runs: {runs_info}")
            
    # Check table 0 (usually title block)
    if doc.tables:
        t0 = doc.tables[0]
        print("  Title table cells:")
        for r_idx, row in enumerate(t0.rows):
            for c_idx, cell in enumerate(row.cells):
                txt = cell.text.strip().replace('\n', ' ')
                fonts = [(p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else "default") for p in cell.paragraphs if p.runs]
                print(f"    Row {r_idx} Col {c_idx}: '{txt[:60]}' | font_sizes={fonts}")

    # 2. Content / TOC check
    print("\n--- 2. TABLE OF CONTENTS (СОДЕРЖАНИЕ) ---")
    toc_lines = []
    in_toc = False
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt == "СОДЕРЖАНИЕ":
            in_toc = True
            continue
        if in_toc:
            if txt == "ВВЕДЕНИЕ" or (p.paragraph_format.page_break_before or "w:br w:type=\"page\"" in p._element.xml):
                if txt != "ВВЕДЕНИЕ	3":  # TOC line
                    break
            toc_lines.append((i, txt))
            if txt.startswith("ПРИЛОЖЕНИЕ"):
                # end of TOC soon
                pass
            if len(toc_lines) > 40:
                break
    print(f"  TOC lines count: {len(toc_lines)}")
    for tl in toc_lines[:35]:
        print(f"    {tl[1]}")

    # 3. Introduction check
    print("\n--- 3. INTRODUCTION (ВВЕДЕНИЕ) ---")
    intro_paragraphs = []
    in_intro = False
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt == "ВВЕДЕНИЕ" and p.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER:
            in_intro = True
            continue
        if in_intro:
            if txt.startswith("1 ") and p.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER:
                break
            intro_paragraphs.append((i, txt))
    print(f"  Intro paragraphs: {len(intro_paragraphs)}")
    intro_words = sum(len(txt.split()) for _, txt in intro_paragraphs)
    print(f"  Intro total words: {intro_words} (~{intro_words/300:.1f} pages)")
    # Check elements:
    intro_text_full = " ".join(txt for _, txt in intro_paragraphs)
    keywords = ["актуальн", "новизн", "значимост", "цель", "объект", "предмет", "гипотез", "задач"]
    for kw in keywords:
        found = kw in intro_text_full.lower()
        print(f"    Keyword '{kw}': {'FOUND' if found else 'MISSING'}")

    # 4. Main Body Structure
    print("\n--- 4. MAIN BODY (ОСНОВНАЯ ЧАСТЬ) ---")
    sections = []
    cur_sec = None
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if p.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER and (txt.startswith("1 ") or txt.startswith("2 ") or txt.startswith("3 ") or txt.startswith("4 ") or txt == "ЗАКЛЮЧЕНИЕ"):
            if cur_sec:
                sections.append(cur_sec)
            cur_sec = {"title": txt, "words": 0, "paragraphs": 0, "subsections": []}
        elif cur_sec:
            if cur_sec["title"] == "ЗАКЛЮЧЕНИЕ" and txt.startswith("СПИСОК"):
                sections.append(cur_sec)
                cur_sec = None
                break
            if txt:
                cur_sec["words"] += len(txt.split())
                cur_sec["paragraphs"] += 1
                if len(txt) > 3 and txt[0].isdigit() and txt[1] == '.':
                    cur_sec["subsections"].append(txt[:60])
    if cur_sec:
        sections.append(cur_sec)

    total_main_words = 0
    for s in sections:
        print(f"  Section: {s['title']}")
        print(f"    Words: {s['words']} (~{s['words']/300:.1f} pages), Paragraphs: {s['paragraphs']}")
        print(f"    Subsections ({len(s['subsections'])}): {s['subsections']}")
        if s['title'] != "ЗАКЛЮЧЕНИЕ":
            total_main_words += s['words']
    print(f"  Total Main Body Words: {total_main_words} (~{total_main_words/300:.1f} pages)")

    # 5. Conclusion check
    print("\n--- 5. CONCLUSION (ЗАКЛЮЧЕНИЕ) ---")
    concl_paras = []
    in_concl = False
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt == "ЗАКЛЮЧЕНИЕ" and p.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER:
            in_concl = True
            continue
        if in_concl:
            if txt.startswith("СПИСОК") and p.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER:
                break
            concl_paras.append((i, txt))
    concl_words = sum(len(txt.split()) for _, txt in concl_paras)
    print(f"  Conclusion paragraphs: {len(concl_paras)}, words: {concl_words} (~{concl_words/300:.1f} pages)")

    # 6. Bibliography check
    print("\n--- 6. BIBLIOGRAPHY (СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ) ---")
    bib_items = []
    in_bib = False
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt == "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ" and p.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER:
            in_bib = True
            continue
        if in_bib:
            if txt.startswith("ПРИЛОЖЕНИЕ"):
                break
            if txt:
                bib_items.append(txt)
    print(f"  Bibliography items count: {len(bib_items)}")
    for b in bib_items:
        print(f"    {b[:70]}")
        
    # Check alphabetical order
    print("\n  Alphabetical order verification:")
    first_letters = []
    for b in bib_items:
        clean = b.lstrip("0123456789. ")
        first_letters.append(clean[:15])
    print(f"  Items start with: {first_letters}")

inspect_contents_and_details("ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA.docx")
inspect_contents_and_details("ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA_ГОТОВО.docx")
