import docx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

def docx_to_markdown(docx_path, md_path):
    doc = docx.Document(docx_path)
    md_lines = []
    
    # Title
    md_lines.append("# ПОЯСНИТЕЛЬНАЯ ЗАПИСКА К ИССЛЕДОВАТЕЛЬСКОМУ ПРОЕКТУ\n")
    md_lines.append("**Тема:** РАЗРАБОТКА МОДУЛЯ ПЕРСОНАЛИЗИРОВАННЫХ РЕКОМЕНДАЦИЙ НА ОСНОВЕ МАШИННОГО ОБУЧЕНИЯ (НА ПРИМЕРЕ МУЗЫКАЛЬНОЙ ИНДУСТРИИ)\n")
    md_lines.append("**Выполнил:** Забинский Артемий, учащийся УО «Национальный детский технопарк»  ")
    md_lines.append("**Руководители проекта:** Сицко В. А., Крюкова Д. С., Парамонов А. И.  ")
    md_lines.append("**Минск, 2026**\n\n---\n")
    
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        
        # Check if heading
        align = p.alignment
        is_bold = bool(p.runs and p.runs[0].bold)
        
        if txt.startswith("МИНИСТЕРСТВО") or txt.startswith("Учреждение") or txt.startswith("Образовательное") or txt == "Исследовательский проект" or txt.startswith("«РАЗРАБОТКА"):
            md_lines.append(f"> **{txt}**\n")
            continue
            
        if txt == "Минск 2026":
            md_lines.append(f"\n> **{txt}**\n\n---\n")
            continue
            
        if txt == "СОДЕРЖАНИЕ":
            md_lines.append(f"\n# {txt}\n")
            continue
            
        if "\t" in txt and any(txt.startswith(k) for k in ["ВВЕДЕНИЕ", "1 ", "1.", "2 ", "2.", "3 ", "3.", "4 ", "4.", "ЗАКЛЮЧЕНИЕ", "СПИСОК", "ПРИЛОЖЕНИЕ"]):
            parts = txt.split("\t")
            title = parts[0].strip()
            page = parts[-1].strip()
            level = 2 if not title[0].isdigit() or '.' not in title[:4] else 3
            prefix = "##" if level == 2 else "###"
            md_lines.append(f"{prefix} {title}\t{page}\n")
            continue
            
        if any(txt.startswith(k) for k in ["ВВЕДЕНИЕ", "1 ТЕОРЕТИЧЕСКИЕ", "2 ПРОЕКТИРОВАНИЕ", "3 ПРОГРАММНАЯ", "4 ЭКСПЕРИМЕНТАЛЬНОЕ", "ЗАКЛЮЧЕНИЕ", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "ПРИЛОЖЕНИЕ А"]):
            md_lines.append(f"\n## {txt}\n")
            continue
            
        if any(txt.startswith(f"{i}.") for i in range(1, 5)):
            md_lines.append(f"\n### {txt}\n")
            continue
            
        if txt.startswith("Рисунок"):
            md_lines.append(f"\n*{txt}*\n")
            continue
            
        if txt.startswith("Таблица"):
            md_lines.append(f"\n**{txt}**\n")
            continue
            
        md_lines.append(f"{txt}\n")
        
    for i, t in enumerate(doc.tables):
        if i == 0: # Title page sign table
            continue
        # Format table as Markdown
        md_lines.append("\n")
        hdr_cells = [c.text.strip().replace("\n", " ") for c in t.rows[0].cells]
        md_lines.append("| " + " | ".join(hdr_cells) + " |")
        md_lines.append("| " + " | ".join(["---"] * len(hdr_cells)) + " |")
        for r in t.rows[1:]:
            row_cells = [c.text.strip().replace("\n", " ") for c in r.cells]
            md_lines.append("| " + " | ".join(row_cells) + " |")
        md_lines.append("\n")
        
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines))
    print(f"Markdown written to: {md_path}")

if __name__ == "__main__":
    docx_to_markdown(
        r"d:\OSPanel\domains\synthesia\ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA.docx",
        r"d:\OSPanel\domains\synthesia\ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA.md"
    )
