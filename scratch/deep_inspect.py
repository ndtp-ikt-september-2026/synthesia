import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

for name, path in [
    ('EMALL', r'C:\Users\zabazaba\Downloads\AyuGram Desktop\РАЗРАБОТКА_МОДУЛЯ_СИНХРОНИЗАЦИИ_ТОВАРОВ_OPENCART_С_МАРКЕТПЛЕЙСОМ.docx'),
    ('CHUGUNOV', r'C:\Users\zabazaba\Downloads\AyuGram Desktop\Чугунов_ПЗ.docx')
]:
    print('='*30, name, '='*30)
    doc = docx.Document(path)
    
    # 1. Figures
    print('--- FIGURE SAMPLES ---')
    fig_count = 0
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.startswith('Рисунок') or t.startswith('Рис'):
            print(f"P{i} [align={p.alignment}, ind={p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0:.1f}mm, sb={p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0}, sa={p.paragraph_format.space_after.pt if p.paragraph_format.space_after else 0}]: '{t}'")
            if p.runs:
                r0 = p.runs[0]
                print(f"   font: {r0.font.name} {r0.font.size.pt if r0.font.size else None}pt bold={r0.bold}")
            fig_count += 1
            if fig_count >= 3:
                break
                
    # 2. Tables
    print('\n--- TABLE SAMPLES ---')
    tbl_count = 0
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.startswith('Таблица') or t.startswith('Табл'):
            print(f"Caption P{i} [align={p.alignment}, ind={p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0:.1f}mm]: '{t}'")
            if p.runs:
                r0 = p.runs[0]
                print(f"   font: {r0.font.name} {r0.font.size.pt if r0.font.size else None}pt bold={r0.bold}")
            tbl_count += 1
            if tbl_count >= 3:
                break
                
    # 3. TOC paragraphs
    print('\n--- TOC SAMPLES ---')
    in_toc = False
    toc_count = 0
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if 'СОДЕРЖАНИЕ' in t:
            in_toc = True
            print(f"TOC header: '{t}'")
            continue
        if in_toc:
            if t.startswith('ВВЕДЕНИЕ') or t.startswith('1 '):
                print(f"TOC item P{i} [align={p.alignment}, ind={p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0:.1f}mm]: '{t}'")
                tabs = [(t.position.mm, t.alignment, t.leader) for t in p.paragraph_format.tab_stops]
                print(f"   tabs: {tabs}")
                toc_count += 1
                if toc_count >= 4:
                    break
            elif t == 'ВВЕДЕНИЕ' and p.paragraph_format.page_break_before:
                break

    # 4. Bibliography
    print('\n--- BIBLIOGRAPHY SAMPLES ---')
    in_bib = False
    bib_count = 0
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if 'СПИСОК' in t and ('ИСТОЧНИК' in t or 'ЛИТЕРАТУР' in t):
            in_bib = True
            print(f"Bib header P{i}: '{t}'")
            continue
        if in_bib:
            if t.startswith('ПРИЛОЖЕНИЕ'):
                break
            if t and (t[0].isdigit() or t.startswith('[')):
                print(f"Bib item P{i} [ind={p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0:.1f}mm, align={p.alignment}]: '{t}'")
                bib_count += 1
                if bib_count >= 3:
                    break

    # 5. Page numbering XML in footer
    print('\n--- FOOTER XML ---')
    sec = doc.sections[0]
    for fp in sec.footer.paragraphs:
        xml = fp._element.xml
        has_page = 'w:fldSimple' in xml or 'w:instrText' in xml
        print(f"Footer text='{fp.text}', align={fp.alignment}, has_page_field={has_page}")
