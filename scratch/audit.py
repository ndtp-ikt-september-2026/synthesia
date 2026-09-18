import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def detailed_audit(filepath):
    print("=" * 60)
    print(f"DETAILED AUDIT FOR: {filepath}")
    print("=" * 60)
    
    doc = docx.Document(filepath)
    
    # 1. Margins & Page Setup
    sec = doc.sections[0]
    w_mm = sec.page_width.mm
    h_mm = sec.page_height.mm
    top_mm = sec.top_margin.mm
    bot_mm = sec.bottom_margin.mm
    left_mm = sec.left_margin.mm
    right_mm = sec.right_margin.mm
    print(f"[PAGE SETUP]")
    print(f"  Format: {'A4' if abs(w_mm-210)<1 and abs(h_mm-297)<1 else f'{w_mm:.1f}x{h_mm:.1f}mm'}")
    print(f"  Margins: Top={top_mm:.1f}mm (req: 20mm), Bottom={bot_mm:.1f}mm (req: 20mm), Left={left_mm:.1f}mm (req: 30mm), Right={right_mm:.1f}mm (req: 10-15mm)")
    print(f"  Different First Page Header/Footer: {sec.different_first_page_header_footer}")
    
    # Check page numbering in footer
    footer = sec.footer
    footer_text = "".join(p.text for p in footer.paragraphs)
    footer_xml = "".join(p._element.xml for p in footer.paragraphs)
    has_page_field = "w:fldSimple" in footer_xml or "w:instrText" in footer_xml or "PAGE" in footer_xml
    footer_alignments = [str(p.alignment) for p in footer.paragraphs]
    print(f"  Footer text: '{footer_text}', Has PAGE field: {has_page_field}, Alignment: {footer_alignments}")
    
    # Check first page footer
    fp_footer = sec.first_page_footer
    fp_footer_text = "".join(p.text for p in fp_footer.paragraphs)
    fp_footer_xml = "".join(p._element.xml for p in fp_footer.paragraphs)
    fp_has_page = "PAGE" in fp_footer_xml
    print(f"  First page footer text: '{fp_footer_text}', Has PAGE field: {fp_has_page}")

    # 2. Paragraph Formatting & Typography
    non_tnr_runs = []
    font_sizes = {}
    alignments = {}
    indents = {}
    line_spacings = {}
    
    for i, p in enumerate(doc.paragraphs):
        align_str = str(p.alignment)
        alignments[align_str] = alignments.get(align_str, 0) + 1
        
        indent = p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0
        indents[round(indent, 1)] = indents.get(round(indent, 1), 0) + 1
        
        ls = p.paragraph_format.line_spacing
        line_spacings[str(ls)] = line_spacings.get(str(ls), 0) + 1
        
        for r in p.runs:
            if r.font.name and r.font.name != "Times New Roman":
                non_tnr_runs.append((i, r.font.name, r.text[:30]))
            sz = round(r.font.size.pt, 1) if r.font.size else "default"
            font_sizes[sz] = font_sizes.get(sz, 0) + 1

    print(f"\n[TYPOGRAPHY & FORMATTING]")
    print(f"  Non-Times New Roman runs count: {len(non_tnr_runs)}")
    if non_tnr_runs:
        print(f"    Sample: {non_tnr_runs[:5]}")
    print(f"  Font sizes distribution: {font_sizes}")
    print(f"  Paragraph alignments distribution: {alignments}")
    print(f"  First line indents (mm) distribution: {indents}")
    print(f"  Line spacing distribution: {line_spacings}")
    
    # 3. Headings Analysis
    print(f"\n[HEADINGS & STRUCTURE]")
    major_sections = ["ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "ПРИЛОЖЕНИЕ"]
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt:
            continue
        
        # Check if section heading
        is_major = any(txt.startswith(kw) for kw in major_sections) or (txt[0].isdigit() and len(txt) > 2 and txt[1] == ' ' and txt.isupper())
        is_sub = (len(txt) > 3 and txt[0].isdigit() and '.' in txt[:4] and not is_major and txt[1] == '.' and (txt[3] == ' ' or (len(txt) > 4 and txt[4] == ' ')))
        
        if is_major:
            # Check font size, bold, all caps, center align, page break before, ending dot
            has_pb = p.paragraph_format.page_break_before or "w:br w:type=\"page\"" in p._element.xml
            # check previous paragraph for page break
            prev_has_pb = False
            if i > 0 and "w:br w:type=\"page\"" in doc.paragraphs[i-1]._element.xml:
                prev_has_pb = True
            
            bold = all(r.bold for r in p.runs if r.text.strip()) if p.runs else False
            sizes = set(round(r.font.size.pt, 1) for r in p.runs if r.font.size)
            has_dot = txt.endswith(".")
            print(f"  SECTION: '{txt[:60]}'")
            print(f"    P#{i} | Align: {p.alignment} (req: CENTER) | Bold: {bold} (req: True) | Sizes: {sizes} (req: 16pt) | Caps: {txt.isupper()} (req: True) | Ends with dot: {has_dot} (req: False) | PageBreak: {has_pb or prev_has_pb} (req: True)")
            
        elif is_sub:
            bold = all(r.bold for r in p.runs if r.text.strip()) if p.runs else False
            sizes = set(round(r.font.size.pt, 1) for r in p.runs if r.font.size)
            has_dot = txt.endswith(".")
            print(f"  SUBSECTION: '{txt[:60]}'")
            print(f"    P#{i} | Align: {p.alignment} (req: CENTER) | Bold: {bold} (req: True) | Sizes: {sizes} (req: 14pt) | Ends with dot: {has_dot} (req: False)")

    # 4. Tables Analysis
    print(f"\n[TABLES] ({len(doc.tables)} tables)")
    for t_idx, t in enumerate(doc.tables):
        # find title paragraph before table
        # find paragraph index
        cell_font_sizes = set()
        cell_fonts = set()
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.size:
                            cell_font_sizes.add(round(r.font.size.pt, 1))
                        if r.font.name:
                            cell_fonts.add(r.font.name)
        print(f"  Table {t_idx}: {len(t.rows)} rows x {len(t.columns)} cols | Cell font sizes: {cell_font_sizes} (req: 10-12pt) | Fonts: {cell_fonts}")

    # 5. Figures Analysis
    print(f"\n[FIGURES / DRAWINGS]")
    fig_paras = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("Рисунок")]
    print(f"  Paragraphs starting with 'Рисунок': {len(fig_paras)}")
    for fp in fig_paras:
        print(f"    {fp[:80]}")

    # 6. Bibliography Analysis
    print(f"\n[BIBLIOGRAPHY / СПИСОК ИСТОЧНИКОВ]")
    bib_started = False
    bib_items = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ" in txt:
            bib_started = True
            continue
        if bib_started:
            if "ПРИЛОЖЕНИЕ" in txt or txt.startswith("1 ") or txt.startswith("2 "):
                break
            if txt and txt[0].isdigit():
                bib_items.append(txt)
    print(f"  Total bibliography items: {len(bib_items)}")
    for b in bib_items[:7]:
        print(f"    {b[:80]}")

detailed_audit("ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA.docx")
detailed_audit("ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA_ГОТОВО.docx")
