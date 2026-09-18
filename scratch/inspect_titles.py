import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

for name, path in [
    ('EMALL', r'C:\Users\zabazaba\Downloads\AyuGram Desktop\РАЗРАБОТКА_МОДУЛЯ_СИНХРОНИЗАЦИИ_ТОВАРОВ_OPENCART_С_МАРКЕТПЛЕЙСОМ.docx'),
    ('CHUGUNOV', r'C:\Users\zabazaba\Downloads\AyuGram Desktop\Чугунов_ПЗ.docx')
]:
    print('='*40, name, '='*40)
    doc = docx.Document(path)
    for i in range(min(32, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        t = p.text.strip()
        if t:
            print(f"P{i} [align={p.alignment}]: '{t}'")
            for r in p.runs:
                if r.text.strip():
                    print(f"   run: text='{r.text.strip()}', sz={r.font.size.pt if r.font.size else None}, bold={r.bold}, font={r.font.name}")
    print('\n')
