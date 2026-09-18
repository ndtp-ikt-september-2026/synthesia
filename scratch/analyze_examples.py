import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def analyze_doc(name, path):
    print("=" * 80)
    print(f"ANALYSIS: {name}")
    print("=" * 80)
    doc = docx.Document(path)
    
    print("\n--- TITLE PAGE ELEMENTS ---")
    for i in range(15):
        p = doc.paragraphs[i]
        t = p.text.strip()
        if t:
            fonts = [(r.text[:25], r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
            print(f"P#{i} [align={p.alignment}]: '{t[:80]}' | {fonts}")
            
    if doc.tables:
        t0 = doc.tables[0]
        print(f"\nTitle Table 0 ({len(t0.rows)}x{len(t0.columns)}):")
        for r_i, row in enumerate(t0.rows):
            for c_i, cell in enumerate(row.cells):
                txt = cell.text.strip().replace('\n', ' ')
                p = cell.paragraphs[0] if cell.paragraphs else None
                runs = [(r.text[:25], r.font.size.pt if r.font.size else None, r.bold) for r in p.runs] if p else []
                print(f"  ({r_i},{c_i}) align={p.alignment if p else None}: '{txt[:80]}' | {runs}")

    print("\n--- HEADINGS (H1 & H2) ---")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if not t: continue
        is_h1 = (t.isupper() and len(t) < 120 and ('ВВЕДЕНИЕ' in t or 'СОДЕРЖАНИЕ' in t or 'ЗАКЛЮЧЕНИЕ' in t or 'СПИСОК' in t or 'ПРИЛОЖЕНИЕ' in t or (t[0].isdigit() and ' ' in t[:3])))
        is_h2 = (len(t) > 3 and t[0].isdigit() and '.' in t[:3] and not is_h1)
        if is_h1:
            runs = [(r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
            pb = p.paragraph_format.page_break_before or "type=\"page\"" in p._element.xml
            print(f"  [H1] P#{i} align={p.alignment} indent={p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0:.1f} pb={pb}: '{t}' | {runs}")
        elif is_h2:
            runs = [(r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
            print(f"  [H2] P#{i} align={p.alignment} indent={p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0:.1f}: '{t}' | {runs}")

    print("\n--- BODY TEXT PARAGRAPHS ---")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if len(t) > 150:
            indent = p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0
            ls = p.paragraph_format.line_spacing
            sb = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0
            sa = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else 0
            runs = [(r.font.name, r.font.size.pt if r.font.size else None) for r in p.runs[:2]]
            print(f"  P#{i} align={p.alignment} indent={indent:.2f}mm ls={ls} sb={sb} sa={sa} | {runs}")
            print(f"    sample: '{t[:100]}...'")
            break

    print("\n--- FIGURES & TABLES CAPTIONS ---")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.startswith("Рисунок") or t.startswith("Рис."):
            runs = [(r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
            indent = p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0
            print(f"  [FIG] P#{i} align={p.alignment} indent={indent:.1f}: '{t}' | {runs}")
        elif t.startswith("Таблица"):
            runs = [(r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
            indent = p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0
            print(f"  [TBL] P#{i} align={p.alignment} indent={indent:.1f}: '{t}' | {runs}")

    print("\n--- TABLE STYLES ---")
    for t_idx, tbl in enumerate(doc.tables):
        if t_idx == 0 and len(tbl.rows) <= 2: continue # skip title
        c0 = tbl.rows[0].cells[0]
        p = c0.paragraphs[0] if c0.paragraphs else None
        runs = [(r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs] if p else []
        print(f"  Table #{t_idx} ({len(tbl.rows)}x{len(tbl.columns)}): align={tbl.alignment}, cell0_align={p.alignment if p else None}, cell0_font={runs}")
        break

    print("\n--- BIBLIOGRAPHY ENTRIES ---")
    in_b = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if "СПИСОК" in t:
            in_b = True
            continue
        if in_b:
            if "ПРИЛОЖЕНИЕ" in t: break
            if t and t[0].isdigit():
                indent = p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0
                runs = [(r.font.name, r.font.size.pt if r.font.size else None) for r in p.runs[:2]]
                print(f"  [BIB] indent={indent:.1f} align={p.alignment}: '{t[:90]}' | {runs}")

analyze_doc("OPENCART С МАРКЕТПЛЕЙСОМ", r'C:\Users\zabazaba\Downloads\AyuGram Desktop\РАЗРАБОТКА_МОДУЛЯ_СИНХРОНИЗАЦИИ_ТОВАРОВ_OPENCART_С_МАРКЕТПЛЕЙСОМ.docx')
analyze_doc("ЧУГУНОВ ПЗ", r'C:\Users\zabazaba\Downloads\AyuGram Desktop\Чугунов_ПЗ.docx')
