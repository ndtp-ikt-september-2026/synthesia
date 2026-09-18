# -*- coding: utf-8 -*-
"""
Модуль стилей и типографики для оформления пояснительной записки
по стандартам УО «Национальный детский технопарк» (Минск 2026)
и правилам ВАК Республики Беларусь.
"""

from docx.shared import Pt, Mm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def setup_document_styles(doc):
    # Настройка параметров страницы (А4, поля: левое 30 мм, правое 10 мм, верхнее 20 мм, нижнее 20 мм)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(10)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)
    
    # Настройка стиля Normal
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.first_line_indent = Mm(12.5)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Для обеспечения шрифта Times New Roman в Word для кириллицы
    rPr = style.element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')
    rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    rPr.append(rFonts)
    
    # Нумерация страниц в колонтитуле (снизу по центру, на 1 странице нет)
    section.different_first_page_header_footer = True
    footer = section.footer
    p_foot = footer.paragraphs[0]
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(0)
    p_foot.paragraph_format.space_after = Pt(0)
    
    r_foot = p_foot.add_run()
    r_foot.font.name = 'Times New Roman'
    r_foot.font.size = Pt(14) # Требование 14 pt
    r_foot.font.color.rgb = RGBColor(0, 0, 0)
    
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r_foot._r.append(fldChar1)
    r_foot._r.append(instrText)
    r_foot._r.append(fldChar2)
    r_foot._r.append(fldChar3)

def add_p(doc, text="", bold_prefix="", italic=False, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=12.5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(first_indent)
    
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = 'Times New Roman'
        rb.font.size = Pt(14)
        rb.font.bold = True
        rb.font.color.rgb = RGBColor(0, 0, 0)
        
    if text:
        rt = p.add_run(text)
        rt.font.name = 'Times New Roman'
        rt.font.size = Pt(14)
        if italic:
            rt.font.italic = True
        rt.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h1(doc, title, page_break=True):
    p = doc.add_paragraph()
    if page_break:
        p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(0)
    
    # Требование Слайд 3: Шрифт 16 pt, полужирный, ПРОПИСНЫЕ буквы, по центру, без точки
    clean_title = title.rstrip('.').upper()
    r = p.add_run(clean_title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h2(doc, title):
    p = doc.add_paragraph()
    # Требование Слайд 3: Шрифт 14 pt, полужирный, ВЫРАВНИВАНИЕ ПО ЦЕНТРУ, без точки в конце
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(0)
    
    clean_title = title.rstrip('.')
    r = p.add_run(clean_title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h3(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(0)
    
    clean_title = title.rstrip('.')
    r = p.add_run(clean_title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_figure(doc, image_path, caption, width_mm=150):
    """
    Вставка иллюстрации по стандарту Слайда 5:
    Все иллюстрации подписываются шрифтом 14 pt: 'Рисунок X.Y. – Название'
    Выравнивание по центру, без точки в конце.
    """
    # Абзац с изображением
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.first_line_indent = Mm(0)
    p_img.paragraph_format.line_spacing = 1.0
    
    r_img = p_img.add_run()
    r_img.add_picture(image_path, width=Mm(width_mm))
    
    # Абзац с подписью
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(12)
    p_cap.paragraph_format.first_line_indent = Mm(0)
    p_cap.paragraph_format.line_spacing = 1.0
    
    r_cap = p_cap.add_run(caption.rstrip('.'))
    r_cap.font.name = 'Times New Roman'
    r_cap.font.size = Pt(14)
    r_cap.font.color.rgb = RGBColor(0, 0, 0)
    return p_img, p_cap

def add_formula(doc, formula_text, formula_num):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(0)
    
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    
    tab_center = OxmlElement('w:tab')
    tab_center.set(qn('w:val'), 'center')
    tab_center.set(qn('w:pos'), '4819')
    tabs.append(tab_center)
    
    tab_right = OxmlElement('w:tab')
    tab_right.set(qn('w:val'), 'right')
    tab_right.set(qn('w:pos'), '9638')
    tabs.append(tab_right)
    
    pPr.append(tabs)
    
    p.add_run("\t")
    rf = p.add_run(formula_text)
    rf.font.name = 'Times New Roman'
    rf.font.size = Pt(14)
    rf.font.italic = True
    
    p.add_run("\t")
    rn = p.add_run(f"({formula_num})")
    rn.font.name = 'Times New Roman'
    rn.font.size = Pt(14)
    return p

def add_table(doc, headers, data, col_widths=None, title="", is_appendix=False):
    """
    Вставка таблицы по стандарту Слайда 5:
    - Нумерация: первая цифра - номер раздела, вторая - порядковый номер (Таблица 1.1)
    - Название: краткое (2-5 слов), шрифт как в основном тексте (14 pt)
    - В ячейках: 10-12 pt (используем 11 pt)
    """
    if title:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_title.paragraph_format.space_before = Pt(10)
        p_title.paragraph_format.space_after = Pt(4)
        p_title.paragraph_format.line_spacing = 1.0
        p_title.paragraph_format.first_line_indent = Mm(12.5)
        
        rt = p_title.add_run(title.rstrip('.'))
        rt.font.name = 'Times New Roman'
        rt.font.size = Pt(14)
        rt.font.color.rgb = RGBColor(0, 0, 0)
        
    num_rows = len(data) + 1
    num_cols = len(headers)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    # Header row
    hdr_row = table.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for c_idx, h_text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F4F8"/>'))
        
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', 120), ('bottom', 120), ('left', 140), ('right', 140)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Mm(0)
        
        r = p.add_run(h_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11) # Требование 10-12 pt
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)
        
    # Data rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            tcPr = cell._tc.get_or_add_tcPr()
            if r_idx % 2 == 1:
                tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAFAFA"/>'))
                
            tcMar = OxmlElement('w:tcMar')
            for m, val_m in [('top', 100), ('bottom', 100), ('left', 140), ('right', 140)]:
                node = OxmlElement(f'w:{m}')
                node.set(qn('w:w'), str(val_m))
                node.set(qn('w:type'), 'dxa')
                tcMar.append(node)
            tcPr.append(tcMar)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 or len(val) > 20 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Mm(0)
            
            r = p.add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11) # Требование 10-12 pt
            r.font.color.rgb = RGBColor(0, 0, 0)
            
    if col_widths:
        for row in table.rows:
            for c_idx, w_mm in enumerate(col_widths):
                if c_idx < len(row.cells):
                    row.cells[c_idx].width = Mm(w_mm)
                    
    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(4)
    p_post.paragraph_format.space_after = Pt(0)
    p_post.paragraph_format.first_line_indent = Mm(0)
    return table
