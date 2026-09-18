import docx
import sys

sys.stdout.reconfigure(encoding="utf-8")

def check_details(path):
    print("=" * 60)
    print("CHECK DETAILS FOR:", path)
    print("=" * 60)
    doc = docx.Document(path)
    
    # Check table titles
    print("\n--- TABLE TITLES IN TEXT ---")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if "Таблица" in txt:
            print(f"  P#{i}: '{txt}' | align={p.alignment} | indent={p.paragraph_format.first_line_indent.mm if p.paragraph_format.first_line_indent else 0}")
            
    # Check drawings or image elements
    print("\n--- DRAWING ELEMENTS IN DOC ---")
    xml = doc._element.xml
    print(f"  w:drawing occurrences: {xml.count('w:drawing')}")
    print(f"  pic:pic occurrences: {xml.count('pic:pic')}")
    print(f"  w:pict occurrences: {xml.count('w:pict')}")
    
    # Check hyphenation setting in doc settings
    print("\n--- HYPHENATION SETTINGS ---")
    try:
        settings_xml = doc.settings._element.xml
        print(f"  autoHyphenation in settings: {'w:autoHyphenation' in settings_xml}")
        print(f"  doNotHyphenate in settings: {'w:doNotHyphenate' in settings_xml}")
    except Exception as e:
        print("  Error reading settings:", e)
        
    # Check paragraphs with manual hyphens or soft hyphens
    soft_hyphens = sum(p.text.count('\xad') for p in doc.paragraphs)
    print(f"  Soft hyphens in text: {soft_hyphens}")

    # Check footer exact XML
    sec = doc.sections[0]
    print("\n--- FOOTER EXACT XML ---")
    for p in sec.footer.paragraphs:
        print(f"  Footer P: align={p.alignment}, xml={p._element.xml}")

    for p in sec.first_page_footer.paragraphs:
        print(f"  First page footer P: align={p.alignment}, xml={p._element.xml}")

check_details("ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA_ГОТОВО.docx")
check_details("ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA.docx")
