import docx
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_doc(path):
    print(f"\n{'='*30} INSPECTING {path} {'='*30}")
    try:
        doc = docx.Document(path)
    except Exception as e:
        print("Error opening docx:", e)
        return

    sec = doc.sections[0]
    print(f"Page size: {sec.page_width.mm:.1f} x {sec.page_height.mm:.1f} mm (A4 is 210 x 297)")
    print(f"Margins: Top={sec.top_margin.mm:.1f}mm, Bottom={sec.bottom_margin.mm:.1f}mm, Left={sec.left_margin.mm:.1f}mm, Right={sec.right_margin.mm:.1f}mm")
    print(f"Different first page header/footer: {sec.different_first_page_header_footer}")

    # Footers
    print("\n-- Footers --")
    for s_idx, s in enumerate(doc.sections):
        f = s.footer
        print(f"Section {s_idx} normal footer has {len(f.paragraphs)} paragraphs:")
        for p in f.paragraphs:
            print(f"   text='{p.text}', align={p.alignment}, xml={p._element.xml[:150]}")
        first_f = s.first_page_footer
        print(f"Section {s_idx} first page footer has {len(first_f.paragraphs)} paragraphs:")
        for p in first_f.paragraphs:
            print(f"   text='{p.text}', align={first_f.paragraphs[0].alignment if first_f.paragraphs else None}")

    total_words = 0
    font_names = set()
    font_sizes = set()
    alignments = set()
    line_spacings = set()
    first_line_indents = set()
    headings = []

    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt:
            continue
        words = txt.split()
        total_words += len(words)
        style_name = p.style.name if p.style else ''
        
        # Check if it looks like a section/subsection heading
        is_upper = txt.isupper() and len(txt) < 100
        has_kw = any(kw in txt.upper() for kw in ['ВВЕДЕНИЕ', 'СОДЕРЖАНИЕ', 'ЗАКЛЮЧЕНИЕ', 'СПИСОК', 'ПРИЛОЖЕНИЕ', 'РАЗДЕЛ'])
        is_heading_style = 'Heading' in style_name or 'Заголовок' in style_name

        if is_upper or has_kw or is_heading_style:
            headings.append((i, style_name, p.alignment, txt[:100]))

        alignments.add(str(p.alignment))
        if p.paragraph_format.first_line_indent:
            first_line_indents.add(round(p.paragraph_format.first_line_indent.mm, 2))
        if p.paragraph_format.line_spacing is not None:
            line_spacings.add(p.paragraph_format.line_spacing)

        for r in p.runs:
            if r.font.name:
                font_names.add(r.font.name)
            if r.font.size:
                font_sizes.add(round(r.font.size.pt, 1))

    print(f"\nTotal paragraphs with text: {len([p for p in doc.paragraphs if p.text.strip()])}")
    print(f"Total words: {total_words}")
    print(f"Estimated pages (~300 words/page): {total_words / 300:.1f} pages (or ~1800 chars with spaces / page: {sum(len(p.text) for p in doc.paragraphs)/1800:.1f} pages)")
    print(f"Total characters with spaces: {sum(len(p.text) for p in doc.paragraphs)}")
    print(f"Fonts detected in runs: {font_names}")
    print(f"Font sizes detected in runs: {font_sizes}")
    print(f"Paragraph alignments: {alignments}")
    print(f"Line spacings: {line_spacings}")
    print(f"First line indents: {first_line_indents}")
    print(f"Tables count: {len(doc.tables)}")

    print("\n-- Headings sample --")
    for h in headings[:40]:
        print(f"  P#{h[0]} [{h[1]}] align={h[2]}: {h[3]}")

    # Check images
    images_count = len(doc.inline_shapes)
    print(f"\nInline shapes / Images: {images_count}")
    # Also check drawing elements in xml
    xml_str = doc._element.xml
    drawing_count = xml_str.count('<w:drawing>')
    pic_count = xml_str.count('<pic:pic')
    print(f"XML drawing tags: {drawing_count}, pic tags: {pic_count}")

if __name__ == '__main__':
    for fname in ['ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA_ГОТОВО.docx', 'ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA.docx', 'ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SOUNDNET.docx']:
        inspect_doc(fname)
