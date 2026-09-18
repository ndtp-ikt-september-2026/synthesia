import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

for fname in ['ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA.docx', 'ПОЯСНИТЕЛЬНАЯ_ЗАПИСКА_SYNTHESIA_ГОТОВО.docx']:
    doc = docx.Document(fname)
    print(f"\n{'='*20} {fname} {'='*20}")
    paragraphs = doc.paragraphs
    full_text = '\n'.join([p.text for p in paragraphs])
    words = len(full_text.split())
    chars = len(full_text)
    print(f"Paragraphs: {len(paragraphs)}, Words: {words}, Characters (with spaces): {chars}")
    
    xml = doc._element.xml
    page_breaks = xml.count('type="page"')
    last_rendered = xml.count('lastRenderedPageBreak')
    print(f"Explicit page breaks: {page_breaks}, Last rendered page breaks: {last_rendered}")
    
    # Check tables
    print(f"Tables count: {len(doc.tables)}")
    for t_idx, t in enumerate(doc.tables):
        print(f"  Table {t_idx}: {len(t.rows)} rows x {len(t.columns)} cols")
        # print text of first row
        header = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
        print(f"    header: {header[:4]}")
        # check font size in cells
        cell_fonts = set()
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.size:
                            cell_fonts.add(r.font.size.pt)
                        if r.font.name:
                            pass
        print(f"    cell font sizes: {cell_fonts}")

