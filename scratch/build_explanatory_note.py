# -*- coding: utf-8 -*-
"""
Генератор полной пояснительной записки проекта SoundNet
в строгом соответствии с ГОСТ 7.32-2001, ГОСТ 2.105-95, ГОСТ 7.1-2003 (ВАК РБ)
и образцами УО «Национальный детский технопарк» (Минск 2026).
"""

import os
import sys
import shutil
from docx import Document
from docx.shared import Pt, Mm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_footer_page_number(section):
    section.different_first_page_header_footer = True
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run()
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r._r.append(fldChar1)
    r._r.append(instrText)
    r._r.append(fldChar2)
    r._r.append(fldChar3)

def add_paragraph(doc, text, bold_prefix="", space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=12.5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(first_indent)
    
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(14)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0, 0, 0)
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_1(doc, title, page_break=True):
    if page_break:
        p = doc.add_paragraph()
        p.paragraph_format.page_break_before = True
    else:
        p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(0)
    
    r = p.add_run(title.upper())
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_2(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(12.5)
    
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_3(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(12.5)
    
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_table_title(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(12.5)
    
    r = p.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_figure_caption(doc, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(0)
    
    r = p.add_run(caption)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_code_line(doc, line, is_first=False, is_last=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4 if is_first else 0)
    p.paragraph_format.space_after = Pt(4 if is_last else 0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Mm(0)
    
    r = p.add_run(line)
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(30, 30, 30)
    return p

print("Core generator functions configured.")
